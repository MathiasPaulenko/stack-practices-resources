"""Runnable examples for the "Read and Write Excel Files in Python" recipe.

Usage:
    python excel_examples.py <example>

Examples:
    sample      Create sample.xlsx with demo data (run this first)
    pandas_read     Read a sheet and all sheets with pandas
    pandas_write    Write a single and multiple sheets with pandas
    openpyxl_format Build a styled report cell by cell
    openpyxl_read   Read cells and ranges with openpyxl
    formulas        Write Excel formulas
    append          Append a sheet to an existing workbook
    large_file      Streaming read with read_only=True
"""

import sys
from pathlib import Path

DATA = {
    "name": ["Alice", "Bob", "Charlie", "Dana", "Eli"],
    "score": [85, 92, 78, 95, 88],
}


def sample() -> None:
    """Create sample.xlsx with two sheets of demo data."""
    import pandas as pd

    df = pd.DataFrame(DATA)
    with pd.ExcelWriter("sample.xlsx") as writer:
        df.to_excel(writer, sheet_name="Sheet1", index=False)
        df.to_excel(writer, sheet_name="Backup", index=False)
    print("Wrote sample.xlsx")


def pandas_read(path: str = "sample.xlsx") -> None:
    import pandas as pd

    df = pd.read_excel(path, sheet_name="Sheet1")
    print(df.head())
    print("Columns:", list(df.columns))

    sheets = pd.read_excel(path, sheet_name=None)
    for name, frame in sheets.items():
        print(f"Sheet: {name}, rows: {len(frame)}")


def pandas_write() -> None:
    import pandas as pd

    df = pd.DataFrame(DATA)
    df.to_excel("output.xlsx", index=False, sheet_name="Results")

    with pd.ExcelWriter("report.xlsx") as writer:
        df.to_excel(writer, sheet_name="Summary", index=False)
        df[df["score"] > 80].to_excel(writer, sheet_name="High Scores", index=False)
    print("Wrote output.xlsx and report.xlsx")


def openpyxl_format() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    headers = ["Name", "Score", "Grade"]
    header_fill = PatternFill(
        start_color="1a56db", end_color="1a56db", fill_type="solid"
    )
    header_font = Font(color="FFFFFF", bold=True)
    thin_bottom = Border(bottom=Side(style="thin", color="1a56db"))

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_bottom
        cell.alignment = Alignment(horizontal="center")

    rows = [("Alice", 85, "B"), ("Bob", 92, "A"), ("Charlie", 78, "C")]
    for row_idx, (name, score, grade) in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=name)
        ws.cell(row=row_idx, column=2, value=score)
        ws.cell(row=row_idx, column=3, value=grade)

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save("formatted_report.xlsx")
    print("Wrote formatted_report.xlsx")


def openpyxl_read(path: str = "sample.xlsx") -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb["Sheet1"]

    for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
        print(row)

    print("A1:", ws["A1"].value)
    wb.close()


def formulas() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active

    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = 30
    ws["A4"] = "=SUM(A1:A3)"
    ws["A5"] = "=AVERAGE(A1:A3)"

    wb.save("formulas.xlsx")
    print("Wrote formulas.xlsx — open it in Excel to see computed values")


def append(path: str = "report.xlsx") -> None:
    import pandas as pd

    if not Path(path).exists():
        pandas_write()

    new_rows = pd.DataFrame({"name": ["Dana"], "score": [95]})
    with pd.ExcelWriter(
        path, mode="a", engine="openpyxl", if_sheet_exists="replace"
    ) as writer:
        new_rows.to_excel(writer, sheet_name="Appended", index=False)
    print(f"Appended sheet to {path}")


def large_file(path: str = "sample.xlsx") -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"]

    count = 0
    for row in ws.iter_rows(values_only=True):
        count += 1
    print(f"Streamed {count} rows from {path}")

    wb.close()


EXAMPLES = {
    "sample": sample,
    "pandas_read": pandas_read,
    "pandas_write": pandas_write,
    "openpyxl_format": openpyxl_format,
    "openpyxl_read": openpyxl_read,
    "formulas": formulas,
    "append": append,
    "large_file": large_file,
}

if __name__ == "__main__":
    if len(sys.argv) != 2 or sys.argv[1] not in EXAMPLES:
        print(__doc__)
        sys.exit(1)
    EXAMPLES[sys.argv[1]]()
