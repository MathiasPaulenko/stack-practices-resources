"""Runnable Excel (.xlsx) examples with openpyxl.

Companion to https://stackpractices.com/recipes/parse-excel-files/

Usage:
    pip install -r requirements.txt
    python excel_parse.py sample     # creates sample.xlsx (run this first)
    python excel_parse.py read       # read rows from sample.xlsx
    python excel_parse.py write      # write output.xlsx
    python excel_parse.py large      # streaming read with read_only=True
"""

import sys
from openpyxl import Workbook, load_workbook


def sample() -> None:
    """Create a sample workbook to parse."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Quantity", "Price"])
    ws.append(["Widget", 100, 19.99])
    ws.append(["Gadget", 40, 4.50])
    ws.append(["Sprocket", 12, 9.75])
    wb.save("sample.xlsx")
    print("Created sample.xlsx")


def read() -> None:
    """Read every row after the header."""
    wb = load_workbook("sample.xlsx")
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        print(row)
    wb.close()


def write() -> None:
    """Write a new workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Quantity", "Price"])
    ws.append(["Widget", 100, 19.99])
    wb.save("output.xlsx")
    print("Created output.xlsx")


def large() -> None:
    """Stream rows without loading the whole workbook into memory."""
    wb = load_workbook("sample.xlsx", read_only=True, data_only=True)
    ws = wb["Sales"]
    for row in ws.iter_rows(values_only=True):
        print(row)
    wb.close()  # required in read_only mode


if __name__ == "__main__":
    commands = {
        "sample": sample,
        "read": read,
        "write": write,
        "large": large,
    }
    cmd = sys.argv[1] if len(sys.argv) > 1 else None
    if cmd in commands:
        commands[cmd]()
    else:
        print(__doc__)
